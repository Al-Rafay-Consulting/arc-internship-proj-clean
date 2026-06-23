import os
import json
import re
import sys
import importlib
from typing import Any, Dict, List
from datetime import datetime
import httpx
from openai import OpenAI
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import time

# Load .env values (if available) before reading API keys.
dotenv_module = importlib.import_module("dotenv") if importlib.util.find_spec("dotenv") else None
if dotenv_module:
    dotenv_module.load_dotenv()


def require_env(name: str) -> str:
    value = os.getenv(name, "").strip()
    if not value:
        if dotenv_module is None:
            print(f"Error: {name} environment variable not set. Install python-dotenv or set the environment variable manually.")
        else:
            print(f"Error: {name} environment variable not set. Add it to your .env file or environment.")
        sys.exit(1)
    return value

# Initialize Qwen (Alibaba DashScope) client via OpenAI-compatible API
qwen_api_key = require_env("DASHSCOPE_API_KEY")

client = OpenAI(
    api_key=qwen_api_key,
    base_url="https://dashscope-intl.aliyuncs.com/compatible-mode/v1",
)

# Model to use — swap to "qwen-max", "qwen-turbo", "qwen-long", etc. as needed
QWEN_MODEL = require_env("QWEN_MODEL")

# Azure AI Search configuration
search_endpoint = require_env("SEARCH_ENDPOINT")
search_key = require_env("SEARCH_KEY")
search_index = require_env("SEARCH_INDEX_NAME")

DOCUMENT_ID = "d0e2081f096848789baf8d677d948edb"
DOCUMENT_FILTER = f"search.ismatch('{DOCUMENT_ID}', 'title')"

SYSTEM_PROMPT = """
# ACTOR
You are a senior commercial real-estate paralegal specializing in lease abstraction. You have read thousands of US commercial leases — office, retail, industrial, and ground leases — and you know how their clauses are organized, how "Basic Lease Provisions" tables summarize key terms, how riders and addenda amend the main body, and where ambiguity typically hides (defined terms, cross-references, exhibits, side letters).
# INTENT
The fields you extract feed a downstream lease abstract that is reviewed by attorneys and used to make portfolio-level financial and legal decisions. False values cause real harm: a wrong rent figure misstates NOI; a wrong renewal-option date causes a missed exercise. 
A null value that signals "not stated" is safer than a confident guess. Every value you return will be cross-checked against the lease, so any fabrication will be caught and will damage the user's trust in this system.
# MISSION
Given a commercial lease document (provided as markdown), extract the requested fields. For each field, return:
  - the value, written in the client house style described below, or null if the lease does not clearly state it; and combined citation as the last independant line, that points to the specific clause where you found the value, so a reviewer can verify it in seconds.

# CITATION COMPLIANCE CONTRACT
- If value is not None, append exactly one citation line as the final line.
- if value is None, do NOT include any citation line.
- Final citation line format must be exactly: [SOURCES: §..., §...]
- Every citation token must start with "§".
- Citation line must contain only comma-separated § references. No prose.
- Length rule: body length limits apply only to the body text. The final citation line is exempt.
- If near length budget, shorten body first. Never drop or truncate the citation line.

# HOUSE STYLE — how the value is written 
The value is legal prose written by a senior paralegal for an attorney to read. These rules apply to every field.
A. DEFINED-TERM REUSE
   When the lease defines a term ("Applicable Laws", "Operating Expenses", "Hazardous Materials", "Tenant Improvements", "Premises",  "Building", "Project"), use the defined term itself. Never expand the underlying definition.
     ✅ "Office and laboratory use in conformity with all Applicable Laws."
     ❌ "Office and laboratory use in conformity with all federal, state, municipal and local laws, codes, ordinances, rules and regulations of Governmental Authorities..."
B. LEGAL VOICE
   For obligations and rights, write in formal legal prose: "Tenant shall...", "Landlord may...", "subject to...", "provided that...", "in the event...".
   For descriptive fields (names, addresses, dates, single amounts), write plain declarative language.
C. NO SUB-HEADED CATEGORIZATION
   A value is one continuous block of legal prose — or, when triggered by rule E below, a bulleted list. Never use labeled sub-sections inside a value:
     ❌ "Permitted use: ... Prohibited activities: ... Compliance: ..."
     ✅ "Tenant may use Hazardous Materials in connection with its business (Lease §21.2), but shall not cause or permit Hazardous Materials to be brought upon the Premises in violation of Applicable Laws (Lease §21.1)."
   If a provision genuinely has multiple distinct mechanisms, separate them with paragraph breaks (\n\n) inside the string — never labels.
D. NO MARKDOWN EXCEPT BULLETS
   No headings (#, ##), no tables (| syntax), no emphasis (**, *, `).
The only formatting permitted is bullets per rule E.
E. BULLETS — CONTENT-SHAPE TRIGGERED
   Render the value as a bulleted list ONLY when ALL THREE triggers are met:
     (i)   The source clause is explicitly enumerated — uses (a)(b)(c), (1)(2)(3), (i)(ii)(iii), or a list-introducing pattern such as "shall not include the following:".
     (ii)  There are FOUR OR MORE discrete items. 1-3 items always stay as inline prose.
     (iii) Items are INDEPENDENT — each stands alone. Comma-separated phrases within one sentence ("structural, roof, HVAC,  plumbing repairs") are NOT independent — keep as prose.
   When triggered:
     - Prefix each line with word style bullets ("• ") e.g.
        •	Bullet point 1
        •	Bullet point 2
     - One item per line, separated by "\n".
     - Close each bullet with a semicolon, except the last (period).
     - No nested sub-bullets. Fold sub-points into one line.
     - Trailing prose context (e.g. "Payable in equal monthly installments") goes on a new line AFTER the last bullet.
# DENSITY — determined by source content, not by field name
  - Single-fact content (a name, address, date, single amount): one short sentence, typically under 250 characters.
  - Single-provision content (one mechanism, right, or obligation): 1-3 sentences, typically 100-500 characters.
  - Multi-provision content (several related mechanisms in one topic): one paragraph or short series, typically 500-2000 characters.
  - Complex multi-mechanism content (with conditions, triggers, carve-outs): multiple paragraphs separated by "\n\n", typically 2000+ characters.
  - Inherently enumerated content (4+ discrete items per rule E): bullets, length scales with item count.
If the lease provides only minimal content, output minimal content —never pad to fill a tier. If the lease provides more material than a
typical tier, exceed it."""


def load_prompts(file_path: str) -> Dict[str, str]:
    with open(file_path, 'r', encoding='utf-8') as f:
        return json.load(f)


def clean_response(text: str) -> str:
    """Remove all [docN] citation markers and clean up whitespace."""
    if not text:
        return text
    cleaned = re.sub(r'\[doc\d+\]', '', text)
    cleaned = re.sub(r'\[Doc\d+\]', '', cleaned)
    cleaned = re.sub(r'\s+\.', '.', cleaned)
    cleaned = re.sub(r'\n{3,}', '\n\n', cleaned)
    return cleaned.strip()


def validate_document_filter() -> None:
    """Verify the Azure Search filter matches at least one document before calling Qwen."""
    if not DOCUMENT_FILTER:
        print("No document filter configured. Skipping filter validation.")
        return
    print(f"Using document filter: {DOCUMENT_FILTER}")

    response = httpx.post(
        f"{search_endpoint}/indexes/{search_index}/docs/search?api-version=2024-07-01",
        headers={
            "api-key": search_key,
            "Content-Type": "application/json",
        },
        json={
            "search": "*",
            "filter": DOCUMENT_FILTER,
            "top": 3,
            "count": True,
        },
        timeout=30.0,
    )
    response.raise_for_status()
    payload = response.json()
    match_count = payload.get("@odata.count", 0)
    if match_count == 0:
        raise RuntimeError(
            "The configured Azure Search filter matched 0 documents. "
            "The field used in DOCUMENT_FILTER is likely wrong for this document id."
        )

    sample_titles = []
    for item in payload.get("value", []):
        if isinstance(item, dict):
            title = item.get("title")
            if title:
                sample_titles.append(title)

    print(f"Filter matched {match_count} document(s).")
    if sample_titles:
        print(f"Sample matched title values: {sample_titles}")


def retrieve_search_context(document_filter: str) -> str:
    """Retrieve document context from Azure Search."""
    if not document_filter:
        return ""

    try:
        response = httpx.post(
            f"{search_endpoint}/indexes/{search_index}/docs/search?api-version=2024-07-01",
            headers={
                "api-key": search_key,
                "Content-Type": "application/json",
            },
            json={
                "search": "*",
                "filter": document_filter,
                "top": 10,
                "select": "*",
            },
            timeout=30.0,
        )
        response.raise_for_status()
        payload = response.json()

        context_parts = []
        for doc in payload.get("value", []):
            if "content" in doc:
                context_parts.append(doc["content"])

        return "\n\n".join(context_parts) if context_parts else ""
    except Exception as e:
        print(f"Warning: Could not retrieve search context: {e}")
        return ""


def call_qwen(user_prompt: str, search_context: str = None, retries: int = 3) -> Dict[str, Any]:
    """Call Qwen API via OpenAI-compatible endpoint with retry logic on failure."""

    # Build the user message, optionally prepending the retrieved document context
    user_content = ""
    if search_context:
        user_content += f"Document Context from Azure Search:\n{search_context}\n\n"
    user_content += f"User Request: {user_prompt}"

    messages = [
        {"role": "system", "content": SYSTEM_PROMPT},
        {"role": "user", "content": user_content},
    ]

    for attempt in range(retries):
        try:
            completion = client.chat.completions.create(
                model=QWEN_MODEL,
                messages=messages,
                max_tokens=6553,
                temperature=0.7,
                top_p=0.95,
            )

            response_content = completion.choices[0].message.content or ""
            cleaned_content = clean_response(response_content)
            finish_reason = completion.choices[0].finish_reason or "stop"
            total_tokens = completion.usage.total_tokens if completion.usage else len(response_content.split())

            return {
                "response": cleaned_content,
                "tokens": total_tokens,
                "finish_reason": finish_reason,
            }

        except Exception as e:
            print(f"    ✗ Attempt {attempt + 1} failed: {str(e)}")
            if attempt < retries - 1:
                wait_time = 5 * (attempt + 1)
                print(f"    Retrying in {wait_time} seconds...")
                time.sleep(wait_time)
            else:
                raise


def create_word_document(results: List[Dict[str, str]], output_path: str) -> None:
    """Create a formatted Word document with results table."""
    doc = Document()

    title = doc.add_heading('Lease Document Extraction Results', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'

    header_cells = table.rows[0].cells
    header_cells[0].text = 'Field'
    header_cells[1].text = 'Value'

    for cell in header_cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.size = Pt(11)

    for result in results:
        row_cells = table.add_row().cells
        row_cells[0].text = result['field']
        row_cells[1].text = result['value']

        for paragraph in row_cells[0].paragraphs:
            for run in paragraph.runs:
                run.font.bold = True

    for row in table.rows:
        row.cells[0].width = Inches(2.0)
        row.cells[1].width = Inches(4.5)

    doc.save(output_path)
    print(f"✓ Word document saved: {output_path}")


def main():
    prompts_file = "default_prompts.json"
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_file = f"output_{timestamp}.docx"

    # Step 1: Load prompts
    print(f"Loading prompts from {prompts_file}...")
    try:
        prompts = load_prompts(prompts_file)
    except FileNotFoundError:
        print(f"Error: {prompts_file} not found.")
        sys.exit(1)
    print(f"Found {len(prompts)} prompts.\n")

    # Step 2: Validate filter
    try:
        validate_document_filter()
    except Exception as exc:
        print(f"Filter validation failed: {exc}")
        sys.exit(1)

    print(f"\nExtracting from document: {DOCUMENT_ID}")
    print(f"Using Qwen model: {QWEN_MODEL}\n")

    # Step 3: Retrieve search context once
    search_context = retrieve_search_context(DOCUMENT_FILTER)

    # Step 4: Process ALL prompts
    results = []
    total_tokens = 0

    for i, (field_name, field_prompt) in enumerate(prompts.items(), 1):
        print(f"[{i}/{len(prompts)}] Processing '{field_name}'...")
        try:
            api_response = call_qwen(field_prompt, search_context)
            total_tokens += api_response.get("tokens", 0)

            print(f"    Response: {api_response['response'][:100]}...")
            results.append({
                "field": field_name,
                "value": api_response["response"]
            })
            print(f"    ✓ Done ({api_response['tokens']} tokens)")
            time.sleep(1)  # 1-second gap between calls to avoid rate limiting

        except Exception as e:
            print(f"    ✗ Failed: {str(e)}")
            results.append({
                "field": field_name,
                "value": f"ERROR: {str(e)}"
            })

    print(f"\nTotal tokens used: {total_tokens}\n")

    # Step 5: Save Word document
    print(f"Creating Word document '{output_file}'...")
    create_word_document(results, output_file)
    print("✓ Batch extraction complete!")


if __name__ == "__main__":
    main()