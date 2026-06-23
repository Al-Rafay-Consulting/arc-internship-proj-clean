import os
import json
import re
import sys
import importlib
from typing import Any, Dict, List
from datetime import datetime
import httpx
from openai import AzureOpenAI
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

dotenv_module = importlib.import_module("dotenv") if importlib.util.find_spec("dotenv") else None
if dotenv_module:
    dotenv_module.load_dotenv()


def require_env(name: str) -> str:
    value = os.getenv(name, "").strip()
    if not value:
        print(f"Error: Missing required environment variable: {name}")
        sys.exit(1)
    return value


# Initialize Azure OpenAI client
endpoint = require_env("ENDPOINT_URL")
deployment = require_env("DEPLOYMENT_NAME")
search_endpoint = require_env("SEARCH_ENDPOINT")
search_key = require_env("SEARCH_KEY")
search_index = require_env("SEARCH_INDEX_NAME")
subscription_key = require_env("OPENAI_API_KEY")

# Document filter configuration
DOCUMENT_ID = "d0e2081f096848789baf8d677d948edb"
DOCUMENT_FILTER = f"search.ismatch('{DOCUMENT_ID}', 'title')"

client = AzureOpenAI(
    azure_endpoint=endpoint,
    api_key=subscription_key,
    api_version="2025-01-01-preview",
)

# System prompt for the model
SYSTEM_PROMPT = """You are a seasoned real estate attorney specializing in the creation and abstraction of lease documents, amendments, and acknowledgment letters. Only respond with information explicitly stated in the provided lease documents. Do not generate, infer, or provide sample content. If information is not found, respond with 'NOT FOUND FOR THIS FIELD'."""

def load_prompts(file_path: str) -> Dict[str, str]:
    """Load prompts from JSON file."""
    with open(file_path, 'r', encoding='utf-8') as f:
        return json.load(f)

def validate_document_filter() -> None:
    """Verify the Azure Search filter matches at least one document before calling OpenAI."""
    if not DOCUMENT_FILTER:
        print("No document filter configured. Skipping filter validation.")
        return
    print(f"Using document filter: {DOCUMENT_FILTER}")

    response = httpx.post(
        f"{search_endpoint}/indexes/{search_index}/docs/search?api-version=2024-07-01",
        headers={
            "api-key": search_key,
            "Content-Type": "application/json",
        },
        json={
            "search": "*",
            "filter": DOCUMENT_FILTER,
            "top": 3,
            "count": True,
        },
        timeout=30.0,
    )
    response.raise_for_status()
    payload = response.json()
    match_count = payload.get("@odata.count", 0)
    if match_count == 0:
        raise RuntimeError(
            "The configured Azure Search filter matched 0 documents. "
            "The field used in DOCUMENT_FILTER is likely wrong for this document id."
        )

    sample_titles = []
    for item in payload.get("value", []):
        if isinstance(item, dict):
            title = item.get("title")
            if title:
                sample_titles.append(title)

    print(f"Filter matched {match_count} document(s).")
    if sample_titles:
        print(f"Sample matched title values: {sample_titles}")

def call_azure_openai(user_prompt: str) -> Dict[str, Any]:
    """Call Azure OpenAI API with the given prompt."""
    messages = [
        {
            "role": "system",
            "content": SYSTEM_PROMPT
        },
        {
            "role": "user",
            "content": user_prompt
        }
    ]
    
    completion = client.chat.completions.create(
        model=deployment,
        messages=messages,
        max_tokens=6553,
        temperature=0.7,
        top_p=0.95,
        frequency_penalty=0,
        presence_penalty=0,
        stop=None,
        stream=False,
        extra_body={
            "data_sources": [{
                "type": "azure_search",
                "parameters": {
                    "endpoint": f"{search_endpoint}",
                    "index_name": f"{search_index}",
                    "semantic_configuration": "default",
                    "query_type": "semantic",
                    "fields_mapping": {},
                    "in_scope": True,
                    "filter": DOCUMENT_FILTER,
                    "strictness": 1,
                    "top_n_documents": 10,
                    "authentication": {
                        "type": "api_key",
                        "key": f"{search_key}"
                    }
                }
            }]
        }
    )
    
    # Extract response data
    response_content = completion.choices[0].message.content
    citations = []
    if hasattr(completion.choices[0].message, 'context') and hasattr(completion.choices[0].message.context, 'citations'):
        citations = completion.choices[0].message.context.citations
    inline_refs = re.findall(r"\[(doc\d+)\]", response_content or "")
    inline_refs = sorted(set(inline_refs))
    
    # Azure OpenAI sometimes returns inline [docN] markers even when context.citations is empty.
    if citations:
        print(f"    Citations: {citations}")
    elif inline_refs:
        print(f"    Inline refs: {inline_refs} (documents were retrieved)")
    else:
        print("    Source refs: None (no retrieved document markers found)")
    
    return {
        "response": response_content,
        "citations": citations,
        "inline_refs": inline_refs,
        "tokens": completion.usage.completion_tokens,
        "finish_reason": completion.choices[0].finish_reason
    }

def create_word_document(results: List[Dict[str, str]], output_path: str) -> None:
    """Create a Word document with results in a formatted table."""
    doc = Document()
    
    # Add title
    title = doc.add_heading('Lease Document Extraction Results', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Add some spacing
    doc.add_paragraph()
    
    # Create table with 2 columns and header row
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    
    # Set header row
    header_cells = table.rows[0].cells
    header_cells[0].text = 'Field'
    header_cells[1].text = 'Value'
    
    # Make header bold
    for cell in header_cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
    
    # Add data rows
    for result in results:
        row_cells = table.add_row().cells
        row_cells[0].text = result['field']
        row_cells[1].text = result['value']
        
        # Make field column bold
        for paragraph in row_cells[0].paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
    
    # Set column widths
    for row in table.rows:
        row.cells[0].width = Inches(2.0)
        row.cells[1].width = Inches(4.0)
    
    # Save document
    doc.save(output_path)
    print(f"✓ Word document created: {output_path}")

def main():
    """Main execution function."""
    prompts_file = "default_prompts.json"
    # Generate timestamped filename to avoid conflicts
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_file = f"output_{timestamp}.docx"
    
    # Load prompts
    print(f"Loading prompts from {prompts_file}...")
    try:
        prompts = load_prompts(prompts_file)
    except FileNotFoundError:
        print(f"Error: {prompts_file} not found.")
        sys.exit(1)

    try:
        validate_document_filter()
    except Exception as exc:
        print(f"Filter validation failed: {exc}")
        sys.exit(1)
    
    print(f"Found {len(prompts)} prompts.\n")
    print(f"Extracting from document: {DOCUMENT_ID}\n")
    
    prompts_to_process = list(prompts.items())
    
    results = []
    total_tokens = 0
    
    # Process each prompt
    for i, (field_name, field_prompt) in enumerate(prompts_to_process, 1):
        print(f"[{i}/{len(prompts_to_process)}] Processing '{field_name}'...")
        print(f"    Prompt: {field_prompt}")
        try:
            api_response = call_azure_openai(field_prompt)
            print(f"    Response: {api_response['response']}")
            results.append({
                "field": field_name,
                "value": api_response["response"]
            })
            total_tokens += api_response["tokens"]
            print(f"    ✓ Response received ({api_response['tokens']} tokens)")
        except Exception as e:
            print(f"    ✗ Error: {str(e)}")
            results.append({
                "field": field_name,
                "value": f"ERROR: {str(e)}"
            })
    
    print(f"\nTotal tokens used: {total_tokens}\n")
    
    # Create Word document
    print(f"Creating Word document '{output_file}'...")
    create_word_document(results, output_file)
    
    print("✓ Batch extraction complete!")

if __name__ == "__main__":
    main()
